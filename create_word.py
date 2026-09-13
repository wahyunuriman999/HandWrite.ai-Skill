from docx import Document
import re

with open('README.md', 'r', encoding='utf-8') as f:
    text = f.read()

# Remove markdown image syntax
text = re.sub(r'!\[.*?\]\(.*?\)', '', text)
# Clean up extra newlines left by image removal
text = re.sub(r'\n{3,}', '\n\n', text)

doc = Document()
doc.add_heading('HandWrite.ai Test Text', 0)
doc.add_paragraph(text.strip())

output_path = 'Teks_Uji_HandWrite.docx'
doc.save(output_path)
print(f"Word document saved to {output_path}")
