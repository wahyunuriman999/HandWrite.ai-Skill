from docx import Document

text = '''
Halo semuanya! Ini adalah dokumen uji coba acak untuk melihat apakah sistem AI HandWrite.ai benar-benar bisa mengubah teks panjang menjadi gambar tulisan tangan dengan gaya yang saya inginkan.

Saya sengaja mengetik paragraf yang cukup panjang agar kita bisa melihat bagaimana sistem melakukan proses "word wrapping" (memotong teks agar tidak keluar batas kertas) serta seberapa rapi jarak antar spasinya.

Jika ini berhasil, maka alat ini akan sangat berguna untuk berbagai macam tugas sekolah, mencatat materi rapat, atau sekadar membuat pesan estetis untuk teman. Mari kita lihat apakah AI ini akan menanyakan 6 pertanyaan gaya tulisan sebelum langsung merender dokumen ini!
'''

doc = Document()
doc.add_heading('Uji Coba HandWrite.ai', 0)
doc.add_paragraph(text.strip())

output_path = 'Dokumen_Test_Asal.docx'
doc.save(output_path)
print(f"Word document saved to {output_path}")
